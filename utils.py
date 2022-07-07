import os
import shutil

import openpyxl
import pandas as pd
from PyQt5.QtWidgets import QFileDialog, QApplication, QMessageBox
from matplotlib import pyplot as plt

pd.options.display.float_format = '{:.2f}'.format


class Utils:
    def __init__(self):
        pass


def check_path(a):
    if not os.path.exists(a):
        os.makedirs(a)
    else:
        print(f'{a} - Path exists')


def get_dataframe(name):
    return pd.read_csv(name)


def getFileNames(self):
    file_filter = 'All Files (*.*);; CSV File (*.csv);; Excel File (*.xlsx *.xls)'
    filedialog = QFileDialog()
    response = filedialog.getOpenFileNames(
        parent=self,
        caption='Select a data file',
        directory=os.getcwd(),
        filter=file_filter,
        initialFilter='Excel File (*.xlsx *.xls)'
    )
    print("File Dialog started")
    # filedialog.exec_()
    return response[0]


def getSaveFileName(self):
    file_filter = 'All Files (*.*);; Data File (*.csv);; Excel File (*.xlsx *.xls)'
    file_dialog = QFileDialog()
    response = file_dialog.getSaveFileName(
        parent=self,
        caption='Select a data file',
        directory='Data File.dat',
        filter=file_filter,
        initialFilter='Excel File (*.xlsx *.xls)'
    )
    print("Save Dialog started")
    return response[0]


def convert_xls2csv(fileNames, output_folder, output_file, isTb=False):
    print("Converting started")
    processed_files = []
    ll = 0
    names = list(map(lambda x: os.path.basename(x), fileNames))
    for name, path in zip(names, fileNames):
        ll += 1
        if not os.path.exists(f'{output_folder}/{name[:10]}v{ll}.csv'):
            data_frame = pd.read_excel(f"{path}", index_col=None)
            # data_frame = data_frame[
            #     ['Journal number', 'Date', 'Ledger account', 'Amount', 'Created date and time']]
            if not isTb:
                data_frame['Ledger account'] = data_frame['Ledger account'].str.slice(0, 6)
            data_frame.to_csv(f'{output_folder}/' + name[:10] + f'v{ll}.csv', encoding='utf-8',
                              index=False)
            processed_files.append(f'{name[:10]}v{ll}.csv')
            print(name + " processed")
        else:
            processed_files = os.listdir(f'{output_folder}/')

    if not os.path.exists(f'{output_file}.csv'):
        # files = [f for f in os.listdir('csv_files') if re.match(r'.*\.cvs', f)]
        # combined_csv = pd.concat([pd.read_csv(f'csv_files/{processed_files}')], join='inner')
        converted = pd.concat([pd.read_csv(f'{output_folder}/{f}') for f in processed_files])
        if not isTb:
            converted.sort_values(['Ledger account'], ascending=True, inplace=True)
        else:
            converted.sort_values(['MainAccount'], ascending=True, inplace=True)
        converted.to_csv(f'{output_file}.csv', index=False, encoding='utf-8-sig')
    else:
        print(f"{output_file}.csv exists")

    return output_file + ".csv"


def save_to_xlsx_file(folder_name, name, dataframe):
    check_path(f'results/{folder_name}')
    dataframe.to_excel(f'results/{folder_name}/{name}.xlsx', index=False, encoding='utf-8')


def plot_data_to_xlsx(file_name, sheet_name, kw):
    xfile = openpyxl.load_workbook(file_name)
    sheet = xfile[sheet_name]
    # print(list(sheet.values))

    # xfile = openpyxl.Workbook()
    # Number of sheets in the workbook (1 sheet in our case)
    # sheet = xfile.worksheets[0]

    fig1, (ax1, ax2) = plt.subplots(figsize=(25, 5), nrows=1, ncols=2)
    ax1.bar([x for x in kw['categories'][:10]], [y for y in kw['values'][:10]])
    ax1.set(xlabel="Words",
            ylabel="Word frequency")

    ax2.bar([x for x in kw['categories'][10:]], [y for y in kw['values'][10:]])
    ax2.set(xlabel="Words",
            ylabel="Word frequency")

    fig1.suptitle("Word Analysis", fontsize=18, fontweight="bold")
    check_path(kw['folder'])
    fig1.savefig(f"{kw['folder']}/{kw['file']}")

    # img = Image('temp.png')
    # img.anchor = 'A1'
    # sheet.add_image(img)
    # # xfile.save('temp.xlsx')
    # xfile.save(file_name)
    #
    # if os.path.exists("temp.png"):
    #     os.remove("temp.png")


def plot_data_to_xlsx2(file_name, sheet_name, kw):
    writer_object = pd.ExcelWriter(file_name,
                                   engine='openpyxl')
    workbook = writer_object.book
    worksheet = writer_object.sheets[sheet_name]
    chart = workbook.add_chart({'type': 'column'})

    # chart = workbook.add_chart({'type': 'bar'})
    chart.add_series({
        # 'categories': f"= {sheet_name} !$A${kw['row']}:$A${kw['row'] + 19}",
        'categories': [sheet_name, kw['row'], 0, kw['row'] + 19, 0],
        # 'values': f"= {sheet_name} !$B${kw['row']}:$B${kw['row'] + 19}",
        'values': [sheet_name, kw['row'], 1, kw['row'] + 19, 1],
    })
    chart.set_title({'name': 'Word analysis'})
    # Add x-axis label
    chart.set_x_axis({'name': 'Words'})
    # Add y-axis label
    chart.set_y_axis({'name': 'Frequency'})
    # Set an Excel chart style.
    chart.set_style(14)
    worksheet.insert_chart(kw['cell_index'], chart)
    writer_object.save()


def output_result_csv(name, dataframe, column_name, merge=True):
    dataframe_main = get_dataframe(name='conv_GL.csv')
    # print(dataframe)
    # print((type(dataframe)))
    dataframe = pd.DataFrame(dataframe, columns=[column_name])
    # print(dataframe)
    # print((type(dataframe)))
    # temp = pd.DataFrame({column_name: temp})
    # temp.reset_index(drop=True, inplace=True)
    # dataframe_main.reset_index(drop=True, inplace=True)

    if merge:
        dataframe = pd.merge(dataframe_main, dataframe, left_index=True, right_index=True, how='inner')
        # dataframe = dataframe_main.merge(dataframe, how='inner')
    else:
        dataframe = dataframe_main.loc[
            dataframe_main['Journal number'].isin(dataframe['Journal number']).to_list()]
    # dataframe.drop_duplicates()
    # dataframe.drop(dataframe.columns[-1], axis=1, inplace=True)
    dataframe.to_csv(f'results/{name}.csv', index=False, encoding='utf-8')
    # dataframe.to_excel(f'results/{name}.xlsx', index=True, encoding='utf-8')


def app_quit(self):
    quit_msg = "Are you sure you want to exit the program?"
    reply = QMessageBox.question(self, 'PyQt5 message', quit_msg,
                                 QMessageBox.Yes | QMessageBox.No, QMessageBox.No)
    if reply == QMessageBox.Yes:
        QApplication.instance().quit()
    else:
        pass
        # QMessageBox.information(self, '', "Nothing Changed")


def run_on_excel(name):
    os.system(f'start excel.exe {os.path.abspath(name)}')


def info_message(self, text):
    # QMessageBox.information(self, '', f'{text}')
    msg = QMessageBox()
    msg.setWindowTitle("PKF JET Helper")
    msg.setText(text)
    msg.exec_()


def copy_file_from_src_to_dest():
    curr_dir = os.getcwd()
    src_path = os.path.join(curr_dir, 'images', 'final_result.xlsx')
    dest_path = os.path.join(curr_dir, 'results')
    shutil.copy(src_path, dest_path)
    return os.path.join(dest_path, 'final_result.xlsx')


# copy_file_from_src_to_dest()
